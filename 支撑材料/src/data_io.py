"""读取并校验题目附件数据。"""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .config import PLATEAU_START_S, RADIUS_MAX_TIME_S
from .models import Environment, RadiusHistory

def read_environment(path: Path) -> Environment:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    time_s = np.asarray([row[0] for row in rows], dtype=float)
    temperature_k = np.asarray([row[1] for row in rows], dtype=float) + 273.15
    moisture = np.asarray([row[2] for row in rows], dtype=float)
    mask = time_s >= PLATEAU_START_S
    return Environment(
        time_s=time_s,
        temperature_k=temperature_k,
        moisture=moisture,
        plateau_temperature_k=float(np.mean(temperature_k[mask])),
        plateau_moisture=float(np.mean(moisture[mask])),
    )


def read_radius(path: Path) -> RadiusHistory:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    time_s = np.asarray([row[0] for row in rows], dtype=float)
    radius_m = np.asarray([row[1] for row in rows], dtype=float) / 100.0
    if time_s[-1] < RADIUS_MAX_TIME_S:
        raise ValueError("附件 2 未覆盖 72 h，不能按既定方案求解问题四")
    if np.any(np.diff(radius_m) > 1e-12):
        raise ValueError("附件 2 半径并非单调不增，请核对数据")
    return RadiusHistory(time_s=time_s, radius_m=radius_m)
