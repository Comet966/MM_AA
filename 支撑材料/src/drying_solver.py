r"""A 题“药材的烘干问题”统一数值求解程序。

建议从支撑材料根目录运行 ``python run_all.py``。程序读取 ``data`` 中的
附件 1/2，求解四问，在 ``result`` 中生成四个 Excel 结果与汇总数据，并
在 ``report/figures`` 中生成图表。
"""

from __future__ import annotations

import argparse
import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.integrate import solve_ivp
from scipy.optimize import brentq
from scipy.sparse import lil_matrix


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
ATTACHMENT_DIR = DATA_DIR
TEMPLATE_DIR = PROJECT_ROOT / "result"
RESULT_DIR = PROJECT_ROOT / "result"
FIGURE_DIR = PROJECT_ROOT / "report" / "figures"

INITIAL_T_K = 28.0 + 273.15
INITIAL_C = 2.55
INITIAL_RADIUS_M = 0.02
HEAT_TRANSFER = 25.0
MASS_TRANSFER = 8.0e-7
PLATEAU_START_S = 7200.0
RADIUS_MAX_TIME_S = 72.0 * 3600.0
FIXED_COMPARISON_MAX_TIME_S = 10.0 * 24.0 * 3600.0
POST_DRY_S = 1800.0
DIFFUSIVITY_SCALE = 1.0

TEMPERATURE_CMAP = LinearSegmentedColormap.from_list(
    "temperature_blue_yellow_red",
    ["#2166ac", "#67a9cf", "#ffffbf", "#fdae61", "#f46d43", "#b2182b"],
)
MOISTURE_CMAP = plt.get_cmap("viridis_r")


@dataclass(frozen=True)
class Environment:
    time_s: np.ndarray
    temperature_k: np.ndarray
    moisture: np.ndarray
    plateau_temperature_k: float
    plateau_moisture: float

    def __call__(self, t: float) -> tuple[float, float]:
        if t < PLATEAU_START_S:
            ta = float(np.interp(t, self.time_s, self.temperature_k))
            ca = float(np.interp(t, self.time_s, self.moisture))
            return ta, ca
        return self.plateau_temperature_k, self.plateau_moisture


@dataclass(frozen=True)
class RadiusHistory:
    time_s: np.ndarray
    radius_m: np.ndarray

    def value_and_rate(self, t: float) -> tuple[float, float]:
        if t < self.time_s[0] - 1e-9 or t > self.time_s[-1] + 1e-9:
            raise ValueError(f"半径查询时间 {t} s 超出附件 2 的 0--72 h 范围")
        tc = float(np.clip(t, self.time_s[0], self.time_s[-1]))
        idx = int(np.searchsorted(self.time_s, tc, side="right") - 1)
        idx = min(max(idx, 0), len(self.time_s) - 2)
        t0, t1 = self.time_s[idx], self.time_s[idx + 1]
        r0, r1 = self.radius_m[idx], self.radius_m[idx + 1]
        rate = float((r1 - r0) / (t1 - t0))
        radius = float(r0 + rate * (tc - t0))
        return radius, rate


@dataclass
class Simulation:
    name: str
    grid: np.ndarray
    solution: object
    moving_radius: bool = False

    @property
    def n(self) -> int:
        return self.grid.size

    def fields(self, time_s: np.ndarray | float) -> tuple[np.ndarray, np.ndarray]:
        values = np.asarray(self.solution.sol(time_s), dtype=float)
        return values[: self.n], values[self.n :]


def read_environment(path: Path) -> Environment:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    time_s = np.asarray([row[0] for row in rows], dtype=float)
    temperature_k = np.asarray([row[1] for row in rows], dtype=float) + 273.15
    moisture = np.asarray([row[2] for row in rows], dtype=float)
    mask = time_s >= PLATEAU_START_S
    return Environment(
        time_s=time_s,
        temperature_k=temperature_k,
        moisture=moisture,
        plateau_temperature_k=float(np.mean(temperature_k[mask])),
        plateau_moisture=float(np.mean(moisture[mask])),
    )


def read_radius(path: Path) -> RadiusHistory:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    time_s = np.asarray([row[0] for row in rows], dtype=float)
    radius_m = np.asarray([row[1] for row in rows], dtype=float) / 100.0
    if time_s[-1] < RADIUS_MAX_TIME_S:
        raise ValueError("附件 2 未覆盖 72 h，不能按既定方案求解问题四")
    if np.any(np.diff(radius_m) > 1e-12):
        raise ValueError("附件 2 半径并非单调不增，请核对数据")
    return RadiusHistory(time_s=time_s, radius_m=radius_m)


def material_properties(
    model: str, temperature_k: np.ndarray, moisture: np.ndarray
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    c = np.maximum(np.asarray(moisture, dtype=float), 1e-10)
    t = np.asarray(temperature_k, dtype=float)
    if model == "appendix2":
        rho = np.full_like(c, 820.0)
        cp = np.full_like(c, 2600.0)
        k = np.full_like(c, 0.36)
        diffusivity = 7.0e-9 * np.exp(-0.89 / c)
    elif model == "appendix3":
        rho = 650.0 + 128.0 * c
        cp = 1450.0 + 2736.0 * c / (c + 1.0)
        k = 0.21 + 0.38 * c / (c + 1.0)
        diffusivity = 2.4e-3 * np.exp(-0.45 / c) * np.exp(-3850.0 / t)
    elif model == "appendix4":
        rho = 760.0 + 90.0 * c
        cp = 1850.0 + 2150.0 * c / (c + 1.0)
        k = 0.12 + 0.20 * c / (c + 1.0)
        diffusivity = 4.2e-4 * np.exp(-0.30 / c) * np.exp(-3850.0 / t)
    else:
        raise ValueError(f"未知物性模型：{model}")
    return rho, cp, k, diffusivity * DIFFUSIVITY_SCALE


def harmonic_mean(values: np.ndarray) -> np.ndarray:
    left, right = values[:-1], values[1:]
    return 2.0 * left * right / np.maximum(left + right, 1e-300)


def coupled_jacobian_sparsity(node_count: int):
    """温度/水分两场、三点空间模板对应的稀疏雅可比结构。"""
    pattern = lil_matrix((2 * node_count, 2 * node_count), dtype=int)
    for equation_field in range(2):
        for variable_field in range(2):
            row_offset = equation_field * node_count
            col_offset = variable_field * node_count
            for i in range(node_count):
                for j in range(max(0, i - 1), min(node_count, i + 2)):
                    pattern[row_offset + i, col_offset + j] = 1
    return pattern.tocsr()


def make_fixed_rhs(
    radius_m: float,
    node_count: int,
    model: str,
    environment: Environment,
) -> tuple[np.ndarray, Callable[[float, np.ndarray], np.ndarray]]:
    r = np.linspace(0.0, radius_m, node_count)
    dr = float(r[1] - r[0])
    r_half = 0.5 * (r[:-1] + r[1:])

    def rhs(t: float, y: np.ndarray) -> np.ndarray:
        temperature = y[:node_count]
        moisture = y[node_count:]
        rho, cp, k, diffusivity = material_properties(model, temperature, moisture)
        k_face = harmonic_mean(k)
        d_face = harmonic_mean(diffusivity)
        ta, ca = environment(t)

        dt = np.empty(node_count, dtype=float)
        dc = np.empty(node_count, dtype=float)
        dt[0] = 4.0 * k_face[0] * (temperature[1] - temperature[0]) / (
            rho[0] * cp[0] * dr**2
        )
        dc[0] = 4.0 * d_face[0] * (moisture[1] - moisture[0]) / dr**2

        heat_out = r_half[1:] * k_face[1:] * np.diff(temperature)[1:] / dr
        heat_in = r_half[:-1] * k_face[:-1] * np.diff(temperature)[:-1] / dr
        mass_out = r_half[1:] * d_face[1:] * np.diff(moisture)[1:] / dr
        mass_in = r_half[:-1] * d_face[:-1] * np.diff(moisture)[:-1] / dr
        dt[1:-1] = (heat_out - heat_in) / (rho[1:-1] * cp[1:-1] * r[1:-1] * dr)
        dc[1:-1] = (mass_out - mass_in) / (r[1:-1] * dr)

        r_minus = radius_m - 0.5 * dr
        shell_measure = radius_m**2 - r_minus**2
        dt[-1] = 2.0 * (
            -radius_m * HEAT_TRANSFER * (temperature[-1] - ta)
            - r_minus * k_face[-1] * (temperature[-1] - temperature[-2]) / dr
        ) / (rho[-1] * cp[-1] * shell_measure)
        dc[-1] = 2.0 * (
            -radius_m * MASS_TRANSFER * (moisture[-1] - ca)
            - r_minus * d_face[-1] * (moisture[-1] - moisture[-2]) / dr
        ) / shell_measure
        return np.concatenate((dt, dc))

    return r, rhs


def make_moving_rhs(
    node_count: int,
    environment: Environment,
    radius_history: RadiusHistory,
    include_geometry_term: bool,
) -> tuple[np.ndarray, Callable[[float, np.ndarray], np.ndarray]]:
    xi = np.linspace(0.0, 1.0, node_count)
    dxi = float(xi[1] - xi[0])
    xi_half = 0.5 * (xi[:-1] + xi[1:])

    def rhs(t: float, y: np.ndarray) -> np.ndarray:
        temperature = y[:node_count]
        moisture = y[node_count:]
        radius, radius_rate = radius_history.value_and_rate(t)
        rho, cp, k, diffusivity = material_properties(
            "appendix4", temperature, moisture
        )
        k_face = harmonic_mean(k)
        d_face = harmonic_mean(diffusivity)
        ta, ca = environment(t)

        dt = np.empty(node_count, dtype=float)
        dc = np.empty(node_count, dtype=float)
        dt[0] = 4.0 * k_face[0] * (temperature[1] - temperature[0]) / (
            rho[0] * cp[0] * radius**2 * dxi**2
        )
        dc[0] = 4.0 * d_face[0] * (moisture[1] - moisture[0]) / (
            radius**2 * dxi**2
        )

        heat_out = xi_half[1:] * k_face[1:] * np.diff(temperature)[1:] / dxi
        heat_in = xi_half[:-1] * k_face[:-1] * np.diff(temperature)[:-1] / dxi
        mass_out = xi_half[1:] * d_face[1:] * np.diff(moisture)[1:] / dxi
        mass_in = xi_half[:-1] * d_face[:-1] * np.diff(moisture)[:-1] / dxi
        dt[1:-1] = (heat_out - heat_in) / (
            rho[1:-1] * cp[1:-1] * radius**2 * xi[1:-1] * dxi
        )
        dc[1:-1] = (mass_out - mass_in) / (
            radius**2 * xi[1:-1] * dxi
        )

        xi_minus = 1.0 - 0.5 * dxi
        shell_measure = 1.0 - xi_minus**2
        dt[-1] = 2.0 * (
            -radius * HEAT_TRANSFER * (temperature[-1] - ta)
            - xi_minus * k_face[-1] * (temperature[-1] - temperature[-2]) / dxi
        ) / (rho[-1] * cp[-1] * radius**2 * shell_measure)
        dc[-1] = 2.0 * (
            -radius * MASS_TRANSFER * (moisture[-1] - ca)
            - xi_minus * d_face[-1] * (moisture[-1] - moisture[-2]) / dxi
        ) / (radius**2 * shell_measure)

        if include_geometry_term:
            factor = radius_rate / radius
            t_xi = np.empty(node_count, dtype=float)
            c_xi = np.empty(node_count, dtype=float)
            t_xi[0] = 0.0
            c_xi[0] = 0.0
            t_xi[1:-1] = (temperature[2:] - temperature[:-2]) / (2.0 * dxi)
            c_xi[1:-1] = (moisture[2:] - moisture[:-2]) / (2.0 * dxi)
            t_xi[-1] = -radius * HEAT_TRANSFER * (temperature[-1] - ta) / k[-1]
            c_xi[-1] = -radius * MASS_TRANSFER * (moisture[-1] - ca) / diffusivity[-1]
            dt += xi * factor * t_xi
            dc += xi * factor * c_xi
        return np.concatenate((dt, dc))

    return xi, rhs


def dry_event(node_count: int) -> Callable[[float, np.ndarray], float]:
    def event(_t: float, y: np.ndarray) -> float:
        return float(np.max(y[node_count:]) - 0.15)

    event.direction = -1
    event.terminal = False
    return event


def solve_fixed(
    name: str,
    end_s: float,
    node_count: int,
    model: str,
    environment: Environment,
    track_dry_event: bool = False,
) -> Simulation:
    grid, rhs = make_fixed_rhs(INITIAL_RADIUS_M, node_count, model, environment)
    y0 = np.concatenate(
        (np.full(node_count, INITIAL_T_K), np.full(node_count, INITIAL_C))
    )
    events = dry_event(node_count) if track_dry_event else None
    solution = solve_ivp(
        rhs,
        (0.0, end_s),
        y0,
        method="BDF",
        rtol=1.0e-6,
        atol=1.0e-8,
        max_step=300.0,
        dense_output=True,
        events=events,
        jac_sparsity=coupled_jacobian_sparsity(node_count),
    )
    if not solution.success:
        raise RuntimeError(f"{name} 求解失败：{solution.message}")
    return Simulation(name=name, grid=grid, solution=solution)


def solve_moving(
    name: str,
    end_s: float,
    node_count: int,
    environment: Environment,
    radius_history: RadiusHistory,
    include_geometry_term: bool,
    track_dry_event: bool = True,
) -> Simulation:
    grid, rhs = make_moving_rhs(
        node_count, environment, radius_history, include_geometry_term
    )
    y0 = np.concatenate(
        (np.full(node_count, INITIAL_T_K), np.full(node_count, INITIAL_C))
    )
    events = dry_event(node_count) if track_dry_event else None
    solution = solve_ivp(
        rhs,
        (0.0, end_s),
        y0,
        method="BDF",
        rtol=1.0e-6,
        atol=1.0e-8,
        max_step=300.0,
        dense_output=True,
        events=events,
        jac_sparsity=coupled_jacobian_sparsity(node_count),
    )
    if not solution.success:
        raise RuntimeError(f"{name} 求解失败：{solution.message}")
    return Simulation(name=name, grid=grid, solution=solution, moving_radius=True)


def event_time(simulation: Simulation) -> float:
    event_arrays = simulation.solution.t_events
    if not event_arrays or len(event_arrays[0]) == 0:
        sample_t = np.linspace(0.0, simulation.solution.t[-1], 2001)
        _, sample_c = simulation.fields(sample_t)
        residual = np.max(sample_c, axis=0) - 0.15
        crossed = np.where((residual[:-1] >= 0.0) & (residual[1:] < 0.0))[0]
        if crossed.size == 0:
            end_h = simulation.solution.t[-1] / 3600.0
            raise RuntimeError(f"{simulation.name} 在 {end_h:g} h 内未达到全域 C<0.15")
        idx = int(crossed[0])

        def scalar_residual(t: float) -> float:
            return float(np.max(simulation.fields(t)[1]) - 0.15)

        return float(brentq(scalar_residual, sample_t[idx], sample_t[idx + 1]))
    return float(event_arrays[0][0])


def sample_fixed(
    simulation: Simulation, times_s: np.ndarray, output_r_m: np.ndarray
) -> tuple[np.ndarray, np.ndarray]:
    temperature, moisture = simulation.fields(times_s)
    sampled_t = np.vstack(
        [np.interp(output_r_m, simulation.grid, temperature[:, j]) for j in range(times_s.size)]
    )
    sampled_c = np.vstack(
        [np.interp(output_r_m, simulation.grid, moisture[:, j]) for j in range(times_s.size)]
    )
    return sampled_t - 273.15, sampled_c


def sample_moving_moisture(
    simulation: Simulation,
    radius_history: RadiusHistory,
    times_s: np.ndarray,
    output_r_m: np.ndarray,
) -> tuple[np.ndarray, np.ndarray]:
    _, moisture = simulation.fields(times_s)
    sampled = np.full((times_s.size, output_r_m.size), np.nan, dtype=float)
    radii = np.empty(times_s.size, dtype=float)
    for j, time_s in enumerate(times_s):
        radius, _ = radius_history.value_and_rate(float(time_s))
        radii[j] = radius
        valid = output_r_m <= radius + 1e-12
        sampled[j, valid] = np.interp(
            output_r_m[valid] / radius, simulation.grid, moisture[:, j]
        )
    return sampled, radii


def reset_sheet(ws) -> None:
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)


def style_header(ws, max_col: int) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1][:max_col]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 13


def write_two_sheet_result(
    path: Path,
    times_s: np.ndarray,
    distances_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture: np.ndarray,
) -> None:
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = "温度"
        wb.create_sheet("水分浓度")
    mapping = {"温度": temperature_c, "水分浓度": moisture}
    for sheet_name, values in mapping.items():
        ws = wb[sheet_name]
        reset_sheet(ws)
        ws.cell(1, 1, "时间\\到药材中心的距离")
        for col, distance in enumerate(distances_cm, start=2):
            ws.cell(1, col, float(round(distance, 1)))
        for row, time_s in enumerate(times_s, start=2):
            ws.cell(row, 1, int(round(float(time_s))))
            for col, value in enumerate(values[row - 2], start=2):
                cell = ws.cell(row, col, float(round(value, 4)))
                cell.number_format = "0.0000"
        style_header(ws, distances_cm.size + 1)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_single_sheet_result(
    path: Path,
    times_s: np.ndarray,
    headers: list[float | str],
    values: np.ndarray,
) -> None:
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = "Sheet1"
    ws = wb.active
    reset_sheet(ws)
    ws.cell(1, 1, "时间\\到药材中心的距离")
    for col, header in enumerate(headers, start=2):
        ws.cell(1, col, header)
    for row, time_s in enumerate(times_s, start=2):
        ws.cell(row, 1, int(round(float(time_s))))
        for col, value in enumerate(values[row - 2], start=2):
            if np.isfinite(value):
                cell = ws.cell(row, col, float(round(value, 4)))
                cell.number_format = "0.0000"
    style_header(ws, len(headers) + 1)
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def representative_table(
    simulation: Simulation, times_h: list[float], distances_cm: list[float]
) -> dict[str, list]:
    times_s = np.asarray(times_h, dtype=float) * 3600.0
    radii_m = np.asarray(distances_cm, dtype=float) / 100.0
    temperature_c, moisture = sample_fixed(simulation, times_s, radii_m)
    return {
        "times_h": times_h,
        "distances_cm": distances_cm,
        "temperature_c": np.round(temperature_c, 4).tolist(),
        "moisture": np.round(moisture, 4).tolist(),
    }


def max_moisture(simulation: Simulation, time_s: float) -> float:
    return float(np.max(simulation.fields(time_s)[1]))


def first_output_time_below(
    simulation: Simulation, step_s: float = 60.0, require_displayed_below: bool = False
) -> tuple[float, float]:
    """返回输出网格上首次严格达标的时刻及未经舍入的最大含水率。"""
    times = np.arange(step_s, simulation.solution.t[-1] + 0.1, step_s)
    maxima = np.max(simulation.fields(times)[1], axis=0)
    eligible = np.round(maxima, 4) < 0.15 if require_displayed_below else maxima < 0.15
    indices = np.flatnonzero(eligible)
    if indices.size == 0:
        qualifier = "四位小数显示" if require_displayed_below else "原始数值"
        raise RuntimeError(f"{simulation.name} 没有在输出网格上以{qualifier}达到 C<0.15")
    index = int(indices[0])
    return float(times[index]), float(maxima[index])


def first_node_output_time_below(
    simulation: Simulation, node_index: int, step_s: float = 60.0
) -> tuple[float, float]:
    """返回指定计算节点在输出网格上首次严格低于0.15的时刻。"""
    times = np.arange(step_s, simulation.solution.t[-1] + 0.1, step_s)
    values = simulation.fields(times)[1][node_index]
    indices = np.flatnonzero(values < 0.15)
    if indices.size == 0:
        raise RuntimeError(f"{simulation.name} 的节点{node_index}在输出网格上未达标")
    index = int(indices[0])
    return float(times[index]), float(values[index])


def grid_difference(
    coarse: Simulation, fine: Simulation, time_s: float, field: str = "moisture"
) -> float:
    coarse_t, coarse_c = coarse.fields(time_s)
    fine_t, fine_c = fine.fields(time_s)
    coarse_values = coarse_c if field == "moisture" else coarse_t
    fine_values = fine_c if field == "moisture" else fine_t
    interpolated = np.interp(coarse.grid, fine.grid, fine_values)
    return float(np.max(np.abs(coarse_values - interpolated)))


def create_figures(
    environment: Environment,
    radius_history: RadiusHistory,
    q1: Simulation,
    q2: Simulation,
    q3: Simulation,
    q4: Simulation,
    appendix4_fixed: Simulation,
    dry3_s: float,
    dry4_s: float,
    safe3_s: float,
    safe4_s: float,
    comparison_times_h: dict[str, float],
) -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.sans-serif": ["Microsoft YaHei", "SimHei"], "axes.unicode_minus": False})

    full_time = np.linspace(0.0, RADIUS_MAX_TIME_S, 721)
    full_env = np.asarray([environment(float(t)) for t in full_time])
    fig, axes = plt.subplots(2, 1, figsize=(7.2, 6.0), sharex=True)
    axes[0].plot(environment.time_s / 3600.0, environment.temperature_k - 273.15, ".", ms=2, label="附件1")
    axes[0].plot(full_time / 3600.0, full_env[:, 0] - 273.15, lw=1.3, label="插值/平台延拓")
    axes[0].set_ylabel("环境温度/°C")
    axes[0].legend()
    axes[1].plot(environment.time_s / 3600.0, environment.moisture, ".", ms=2)
    axes[1].plot(full_time / 3600.0, full_env[:, 1], lw=1.3)
    axes[1].set_xlabel("时间/h")
    axes[1].set_ylabel("环境水分浓度/(kg/kg)")
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "environment_extension.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.0, 3.8))
    for t_s in [100, 600, 1200, 1800]:
        tc, cc = sample_fixed(q1, np.asarray([t_s]), q1.grid)
        axes[0].plot(q1.grid * 100, tc[0], label=f"{t_s}s")
        axes[1].plot(q1.grid * 100, cc[0], label=f"{t_s}s")
    axes[0].set(xlabel="到中心距离/cm", ylabel="温度/°C")
    axes[1].set(xlabel="到中心距离/cm", ylabel="水分浓度/(kg/kg)")
    axes[1].legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "question1_profiles.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.0, 3.8))
    for t_h in [0.5, 1.0, 2.0, 3.0]:
        tc, cc = sample_fixed(q2, np.asarray([t_h * 3600]), q2.grid)
        axes[0].plot(q2.grid * 100, tc[0], label=f"{t_h:g}h")
        axes[1].plot(q2.grid * 100, cc[0], label=f"{t_h:g}h")
    axes[0].set(xlabel="到中心距离/cm", ylabel="温度/°C")
    axes[1].set(xlabel="到中心距离/cm", ylabel="水分浓度/(kg/kg)")
    axes[1].legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "question2_profiles.png", dpi=180)
    plt.close(fig)

    def fixed_heatmaps(simulation: Simulation, end_s: float, filename: str, title: str) -> None:
        plot_times = np.linspace(0.0, end_s, 401)
        plot_r_m = np.linspace(0.0, INITIAL_RADIUS_M, 241)
        temperature_c, moisture = sample_fixed(simulation, plot_times, plot_r_m)
        time_scale = 3600.0 if end_s > 7200.0 else 60.0
        time_label = "时间/h" if time_scale == 3600.0 else "时间/min"
        fig, axes = plt.subplots(1, 2, figsize=(9.4, 4.0), sharey=True)
        image_t = axes[0].pcolormesh(
            plot_r_m * 100.0, plot_times / time_scale, temperature_c,
            shading="auto", cmap=TEMPERATURE_CMAP
        )
        image_c = axes[1].pcolormesh(
            plot_r_m * 100.0, plot_times / time_scale, moisture,
            shading="auto", cmap=MOISTURE_CMAP
        )
        axes[0].set(xlabel="到中心距离/cm", ylabel=time_label, title="温度/°C")
        axes[1].set(xlabel="到中心距离/cm", title="水分浓度/(kg/kg)")
        fig.colorbar(image_t, ax=axes[0], pad=0.02)
        fig.colorbar(image_c, ax=axes[1], pad=0.02)
        fig.suptitle(title)
        fig.tight_layout()
        fig.savefig(FIGURE_DIR / filename, dpi=180)
        plt.close(fig)

    fixed_heatmaps(q1, 1800.0, "question1_heatmaps.png", "问题一完整时空演化")
    fixed_heatmaps(q2, 10800.0, "question2_heatmaps.png", "问题二完整时空演化")
    fixed_heatmaps(q3, safe3_s + POST_DRY_S, "question3_heatmaps.png", "问题三至达标后0.5 h的时空演化")

    q4_plot_times = np.linspace(0.0, safe4_s + POST_DRY_S, 401)
    q4_plot_r_m = np.linspace(0.0, INITIAL_RADIUS_M, 241)
    q4_temperature, q4_moisture = q4.fields(q4_plot_times)
    q4_t_map = np.full((q4_plot_times.size, q4_plot_r_m.size), np.nan)
    q4_c_map = np.full_like(q4_t_map, np.nan)
    for j, time_s in enumerate(q4_plot_times):
        radius, _ = radius_history.value_and_rate(float(time_s))
        valid = q4_plot_r_m <= radius + 1e-12
        xi_query = q4_plot_r_m[valid] / radius
        q4_t_map[j, valid] = np.interp(xi_query, q4.grid, q4_temperature[:, j]) - 273.15
        q4_c_map[j, valid] = np.interp(xi_query, q4.grid, q4_moisture[:, j])
    cmap_t = TEMPERATURE_CMAP.copy()
    cmap_c = MOISTURE_CMAP.copy()
    cmap_t.set_bad("white")
    cmap_c.set_bad("white")
    fig, axes = plt.subplots(1, 2, figsize=(9.6, 4.2), sharey=True, constrained_layout=True)
    image_t = axes[0].pcolormesh(
        q4_plot_r_m * 100.0, q4_plot_times / 3600.0, q4_t_map,
        shading="auto", cmap=cmap_t
    )
    image_c = axes[1].pcolormesh(
        q4_plot_r_m * 100.0, q4_plot_times / 3600.0, q4_c_map,
        shading="auto", cmap=cmap_c
    )
    radius_curve = np.asarray([radius_history.value_and_rate(float(t))[0] for t in q4_plot_times])
    for axis in axes:
        axis.plot(radius_curve * 100.0, q4_plot_times / 3600.0, color="white", lw=1.2, ls="--")
        axis.set_xlabel("到中心距离/cm")
    axes[0].set(ylabel="时间/h", title="温度/°C")
    axes[1].set_title("水分浓度/(kg/kg)")
    fig.colorbar(image_t, ax=axes[0], pad=0.02)
    fig.colorbar(image_c, ax=axes[1], pad=0.02)
    fig.suptitle("问题四移动物理域时空演化（白色区域为药材外部）")
    fig.savefig(FIGURE_DIR / "question4_heatmaps.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    check_time3 = np.linspace(0.0, min(q3.solution.t[-1], dry3_s + POST_DRY_S), 800)
    check_time4 = np.linspace(0.0, min(q4.solution.t[-1], dry4_s + POST_DRY_S), 500)
    max3 = np.max(q3.fields(check_time3)[1], axis=0)
    max4 = np.max(q4.fields(check_time4)[1], axis=0)
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    ax.plot(check_time3 / 3600.0, max3, label="问题3：固定半径/附录3")
    ax.plot(check_time4 / 3600.0, max4, label="问题4：动态半径/附录4")
    ax.axhline(0.15, color="black", ls="--", lw=1, label="阈值0.15")
    safe_values = [max_moisture(q3, safe3_s), max_moisture(q4, safe4_s)]
    ax.scatter([safe3_s / 3600, safe4_s / 3600], safe_values, s=24, label="60 s网格首个严格达标点")
    ax.set(xlabel="时间/h", ylabel="全域最大水分浓度/(kg/kg)")
    ax.set_ylim(bottom=0)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "drying_threshold.png", dpi=180)
    plt.close(fig)

    fixed_end_s = event_time(appendix4_fixed)
    fixed_times = np.linspace(0.0, min(appendix4_fixed.solution.t[-1], fixed_end_s + POST_DRY_S), 900)
    moving_times = np.linspace(0.0, min(q4.solution.t[-1], safe4_s + POST_DRY_S), 500)
    fixed_max = np.max(appendix4_fixed.fields(fixed_times)[1], axis=0)
    moving_max = np.max(q4.fields(moving_times)[1], axis=0)
    fig, ax = plt.subplots(figsize=(7.4, 4.3))
    ax.plot(fixed_times / 3600.0, fixed_max, lw=1.5, label="附录4物性：固定半径2 cm")
    ax.plot(moving_times / 3600.0, moving_max, lw=1.5, label="附录4物性：附件2动态半径")
    ax.axhline(0.15, color="black", ls="--", lw=1, label="阈值0.15")
    ax.scatter(
        [fixed_end_s / 3600.0, dry4_s / 3600.0], [0.15, 0.15],
        color=["#4C78A8", "#F58518"], zorder=3
    )
    ax.annotate(f"{fixed_end_s / 3600.0:.2f} h", (fixed_end_s / 3600.0, 0.15), xytext=(-38, 12), textcoords="offset points")
    ax.annotate(f"{dry4_s / 3600.0:.2f} h", (dry4_s / 3600.0, 0.15), xytext=(6, 12), textcoords="offset points")
    ax.set(xlabel="时间/h", ylabel="全域最大水分浓度/(kg/kg)", ylim=(0, None))
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "shrinkage_effect.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.0, 3.8))
    axes[0].plot(radius_history.time_s / 3600.0, radius_history.radius_m * 100.0)
    axes[0].set(xlabel="时间/h", ylabel="半径/cm")
    names = list(comparison_times_h)
    vals = [comparison_times_h[name] for name in names]
    axes[1].bar(range(len(vals)), vals, color=["#4C78A8", "#F58518", "#54A24B", "#B279A2"][: len(vals)])
    axes[1].set_xticks(range(len(vals)), names, rotation=18, ha="right", fontsize=8)
    axes[1].set_ylabel("烘干达标时间/h")
    for idx, value in enumerate(vals):
        axes[1].text(idx, value, f"{value:.2f}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    fig.savefig(FIGURE_DIR / "radius_and_comparison.png", dpi=180)
    plt.close(fig)


def run_all(write_excel: bool = True) -> dict:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    environment = read_environment(ATTACHMENT_DIR / "附件1.xlsx")
    radius_history = read_radius(ATTACHMENT_DIR / "附件2.xlsx")

    print("[1/8] 求解问题一（161 节点）", flush=True)
    q1 = solve_fixed("问题一", 1800.0, 161, "appendix2", environment)
    print("[2/8] 求解问题二（201 节点）", flush=True)
    q2 = solve_fixed("问题二", 10800.0, 201, "appendix3", environment)
    print("[3/8] 求解问题三（801 节点，积分至 72 h）", flush=True)
    q3 = solve_fixed("问题三", RADIUS_MAX_TIME_S, 801, "appendix3", environment, True)
    dry3_s = event_time(q3)
    safe3_s, safe3_c = first_output_time_below(q3)
    visible3_s, visible3_c = first_output_time_below(q3, require_displayed_below=True)
    print(
        f"      问题三连续临界：{dry3_s / 3600:.6f} h；"
        f"60 s网格首个严格达标：{safe3_s:.0f} s",
        flush=True,
    )
    print("[4/8] 求解问题四（601 节点，含几何项）", flush=True)
    q4 = solve_moving("问题四", RADIUS_MAX_TIME_S, 601, environment, radius_history, True, True)
    dry4_s = event_time(q4)
    safe4_s, safe4_c = first_output_time_below(q4)
    visible4_s, visible4_c = first_output_time_below(q4, require_displayed_below=True)
    print(
        f"      问题四连续临界：{dry4_s / 3600:.6f} h；"
        f"60 s网格首个严格达标：{safe4_s:.0f} s",
        flush=True,
    )

    print("[5/8] 求解网格与模型敏感性对照", flush=True)
    q1_coarse = solve_fixed("问题一粗网格", 1800.0, 81, "appendix2", environment)
    q1_fine = solve_fixed("问题一细网格", 1800.0, 321, "appendix2", environment)
    q2_coarse = solve_fixed("问题二粗网格", 10800.0, 101, "appendix3", environment)
    q2_fine = solve_fixed("问题二细网格", 10800.0, 401, "appendix3", environment)
    q3_coarse = solve_fixed("问题三粗网格", RADIUS_MAX_TIME_S, 401, "appendix3", environment, True)
    q3_fine = solve_fixed("问题三细网格", RADIUS_MAX_TIME_S, 1201, "appendix3", environment, True)
    q4_coarse = solve_moving("问题四粗网格", RADIUS_MAX_TIME_S, 321, environment, radius_history, True, True)
    q4_fine = solve_moving("问题四细网格", RADIUS_MAX_TIME_S, 801, environment, radius_history, True, True)
    q4_no_geometry = solve_moving("问题四无几何项", RADIUS_MAX_TIME_S, 601, environment, radius_history, False, True)
    appendix4_fixed = solve_fixed("附录4固定半径", FIXED_COMPARISON_MAX_TIME_S, 801, "appendix4", environment, True)
    dry3_coarse_s = event_time(q3_coarse)
    dry3_fine_s = event_time(q3_fine)
    dry4_coarse_s = event_time(q4_coarse)
    dry4_fine_s = event_time(q4_fine)
    dry4_no_geometry_s = event_time(q4_no_geometry)
    dry_appendix4_fixed_s = event_time(appendix4_fixed)
    safe_appendix4_fixed_s, safe_appendix4_fixed_c = first_output_time_below(appendix4_fixed)
    q3_surface_safe_s, q3_surface_safe_c = first_node_output_time_below(q3, -1)
    q4_radius_at_safe_m, _ = radius_history.value_and_rate(safe4_s)

    output_r_cm = np.round(np.arange(0.0, 2.0 + 0.05, 0.1), 1)
    output_r_m = output_r_cm / 100.0
    if write_excel:
        print("[6/8] 填写四个 Excel 结果模板", flush=True)
        q1_times = np.arange(1.0, 1800.0 + 1.0, 1.0)
        q1_t, q1_c = sample_fixed(q1, q1_times, output_r_m)
        write_two_sheet_result(TEMPLATE_DIR / "result1.xlsx", q1_times, output_r_cm, q1_t, q1_c)

        q2_times = np.arange(1.0, 10800.0 + 1.0, 1.0)
        q2_t, q2_c = sample_fixed(q2, q2_times, output_r_m)
        write_two_sheet_result(TEMPLATE_DIR / "result2.xlsx", q2_times, output_r_cm, q2_t, q2_c)

        q3_end_s = min(RADIUS_MAX_TIME_S, safe3_s + POST_DRY_S)
        q3_times = np.arange(60.0, q3_end_s + 0.1, 60.0)
        _, q3_c = sample_fixed(q3, q3_times, output_r_m)
        write_single_sheet_result(
            TEMPLATE_DIR / "result3.xlsx", q3_times, output_r_cm.tolist(), q3_c
        )

        q4_end_s = min(RADIUS_MAX_TIME_S, safe4_s + POST_DRY_S)
        q4_times = np.arange(60.0, q4_end_s + 0.1, 60.0)
        q4_c, _ = sample_moving_moisture(q4, radius_history, q4_times, output_r_m)
        q4_surface = q4.fields(q4_times)[1][-1, :]
        q4_values = np.column_stack((q4_c, q4_surface))
        q4_headers: list[float | str] = output_r_cm.tolist() + ["药材表面"]
        write_single_sheet_result(
            TEMPLATE_DIR / "result4.xlsx", q4_times, q4_headers, q4_values
        )

    print("[7/8] 汇总代表点、误差与可行性检查", flush=True)
    q1_paper_times = np.asarray([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
    paper_r_cm = np.asarray([0, 0.5, 1.0, 1.5, 2.0], dtype=float)
    q1_pt, q1_pc = sample_fixed(q1, q1_paper_times, paper_r_cm / 100.0)
    q2_paper = representative_table(q2, [0.5, 1, 1.5, 2, 2.5, 3], paper_r_cm.tolist())

    q3_paper_hours = np.append(
        np.arange(6.0, math.floor(safe3_s / 3600.0 / 6.0) * 6.0 + 0.1, 6.0),
        safe3_s / 3600.0,
    )
    _, q3_paper_c = sample_fixed(q3, q3_paper_hours * 3600.0, paper_r_cm / 100.0)
    q4_fixed_r_cm = np.asarray([0.0, 0.5, 1.0, 1.5], dtype=float)
    q4_paper_hours = np.append(
        np.arange(6.0, math.floor(safe4_s / 3600.0 / 6.0) * 6.0 + 0.1, 6.0),
        safe4_s / 3600.0,
    )
    q4_paper_c, q4_paper_radii = sample_moving_moisture(
        q4, radius_history, q4_paper_hours * 3600.0, q4_fixed_r_cm / 100.0
    )
    q4_paper_surface = q4.fields(q4_paper_hours * 3600.0)[1][-1, :]

    def nullable_rows(array: np.ndarray) -> list[list[float | None]]:
        return [
            [None if not np.isfinite(value) else round(float(value), 4) for value in row]
            for row in array
        ]

    validation_delta_s = 60.0
    dry3_before = max_moisture(q3, max(0.0, dry3_s - validation_delta_s))
    dry3_after = max_moisture(q3, min(RADIUS_MAX_TIME_S, dry3_s + POST_DRY_S))
    dry4_before = max_moisture(q4, max(0.0, dry4_s - validation_delta_s))
    dry4_after = max_moisture(q4, min(RADIUS_MAX_TIME_S, dry4_s + POST_DRY_S))
    safe3_previous = max_moisture(q3, safe3_s - 60.0)
    safe4_previous = max_moisture(q4, safe4_s - 60.0)

    comparison_times_h = {
        "附录3固定半径": dry3_s / 3600.0,
        "附录4固定半径": dry_appendix4_fixed_s / 3600.0,
        "附录4动态半径": dry4_s / 3600.0,
        "动态半径无几何项": dry4_no_geometry_s / 3600.0,
    }
    summary = {
        "environment": {
            "plateau_start_s": PLATEAU_START_S,
            "plateau_temperature_c": environment.plateau_temperature_k - 273.15,
            "plateau_moisture": environment.plateau_moisture,
        },
        "radius": {
            "initial_cm": radius_history.radius_m[0] * 100.0,
            "at_72h_cm": radius_history.radius_m[-1] * 100.0,
        },
        "drying_time": {
            "question3_continuous_threshold_s": dry3_s,
            "question3_continuous_threshold_h": dry3_s / 3600.0,
            "question3_safe_60s_output_s": safe3_s,
            "question3_safe_60s_output_h": safe3_s / 3600.0,
            "question3_first_displayed_below_s": visible3_s,
            "question3_first_displayed_below_h": visible3_s / 3600.0,
            "question4_continuous_threshold_s": dry4_s,
            "question4_continuous_threshold_h": dry4_s / 3600.0,
            "question4_safe_60s_output_s": safe4_s,
            "question4_safe_60s_output_h": safe4_s / 3600.0,
            "question4_first_displayed_below_s": visible4_s,
            "question4_first_displayed_below_h": visible4_s / 3600.0,
            "appendix4_fixed_h": dry_appendix4_fixed_s / 3600.0,
            "question4_without_geometry_term_h": dry4_no_geometry_s / 3600.0,
        },
        "threshold_check": {
            "question3_safe_60s_output_raw_max": safe3_c,
            "question3_safe_60s_output_rounded_4dp": round(safe3_c, 4),
            "question3_previous_60s_output_raw_max": safe3_previous,
            "question3_first_displayed_below_raw_max": visible3_c,
            "question3_60s_before": dry3_before,
            "question3_30min_after": dry3_after,
            "question4_safe_60s_output_raw_max": safe4_c,
            "question4_safe_60s_output_rounded_4dp": round(safe4_c, 4),
            "question4_previous_60s_output_raw_max": safe4_previous,
            "question4_first_displayed_below_raw_max": visible4_c,
            "question4_60s_before": dry4_before,
            "question4_30min_after": dry4_after,
        },
        "abstract_metrics": {
            "question1_at_1800s": {
                "center_temperature_c": float(q1_pt[-1, 0]),
                "surface_temperature_c": float(q1_pt[-1, -1]),
                "center_moisture": float(q1_pc[-1, 0]),
                "surface_moisture": float(q1_pc[-1, -1]),
            },
            "question2_at_3h": {
                "center_temperature_c": float(q2_paper["temperature_c"][-1][0]),
                "surface_temperature_c": float(q2_paper["temperature_c"][-1][-1]),
                "center_moisture": float(q2_paper["moisture"][-1][0]),
                "surface_moisture": float(q2_paper["moisture"][-1][-1]),
            },
            "question3_surface_safe_60s_s": q3_surface_safe_s,
            "question3_surface_safe_60s_h": q3_surface_safe_s / 3600.0,
            "question3_surface_safe_raw_moisture": q3_surface_safe_c,
            "question3_center_to_surface_drying_time_ratio": safe3_s / q3_surface_safe_s,
            "question3_center_minus_surface_drying_time_h": (safe3_s - q3_surface_safe_s) / 3600.0,
            "question3_center_to_surface_moisture_ratio_at_finish": safe3_c / float(q3.fields(safe3_s)[1][-1]),
            "question4_radius_at_finish_cm": q4_radius_at_safe_m * 100.0,
            "appendix4_fixed_safe_60s_s": safe_appendix4_fixed_s,
            "appendix4_fixed_safe_60s_h": safe_appendix4_fixed_s / 3600.0,
            "appendix4_fixed_safe_raw_moisture": safe_appendix4_fixed_c,
            "question4_time_reduction_vs_fixed_h": (safe_appendix4_fixed_s - safe4_s) / 3600.0,
            "question4_time_reduction_vs_fixed_percent": (safe_appendix4_fixed_s - safe4_s) / safe_appendix4_fixed_s * 100.0,
            "question4_drying_speed_ratio_vs_fixed": safe_appendix4_fixed_s / safe4_s,
        },
        "grid_convergence": {
            "question1_max_temperature_difference_81_vs_321_at_1800s_K": grid_difference(q1_coarse, q1_fine, 1800.0, "temperature"),
            "question1_max_moisture_difference_81_vs_321_at_1800s": grid_difference(q1_coarse, q1_fine, 1800.0, "moisture"),
            "question1_max_temperature_difference_161_vs_321_at_1800s_K": grid_difference(q1, q1_fine, 1800.0, "temperature"),
            "question1_max_moisture_difference_161_vs_321_at_1800s": grid_difference(q1, q1_fine, 1800.0, "moisture"),
            "question2_max_temperature_difference_101_vs_401_at_3h_K": grid_difference(q2_coarse, q2_fine, 10800.0, "temperature"),
            "question2_max_moisture_difference_101_vs_401_at_3h": grid_difference(q2_coarse, q2_fine, 10800.0, "moisture"),
            "question2_max_temperature_difference_201_vs_401_at_3h_K": grid_difference(q2, q2_fine, 10800.0, "temperature"),
            "question2_max_moisture_difference_201_vs_401_at_3h": grid_difference(q2, q2_fine, 10800.0, "moisture"),
            "question3_dry_time_h_401_nodes": dry3_coarse_s / 3600.0,
            "question3_dry_time_h_801_nodes": dry3_s / 3600.0,
            "question3_dry_time_h_1201_nodes": dry3_fine_s / 3600.0,
            "question4_dry_time_h_321_nodes": dry4_coarse_s / 3600.0,
            "question4_dry_time_h_601_nodes": dry4_s / 3600.0,
            "question4_dry_time_h_801_nodes": dry4_fine_s / 3600.0,
        },
        "paper_tables": {
            "question1": {
                "times_s": q1_paper_times.astype(int).tolist(),
                "distances_cm": paper_r_cm.tolist(),
                "temperature_c": np.round(q1_pt, 4).tolist(),
                "moisture": np.round(q1_pc, 4).tolist(),
            },
            "question2": q2_paper,
            "question3": {
                "times_h": np.round(q3_paper_hours, 6).tolist(),
                "distances_cm": paper_r_cm.tolist(),
                "moisture": np.round(q3_paper_c, 4).tolist(),
            },
            "question4": {
                "times_h": np.round(q4_paper_hours, 6).tolist(),
                "fixed_distances_cm": q4_fixed_r_cm.tolist(),
                "radius_cm": np.round(q4_paper_radii * 100.0, 4).tolist(),
                "moisture_at_fixed_distances": nullable_rows(q4_paper_c),
                "surface_moisture": np.round(q4_paper_surface, 4).tolist(),
            },
        },
        "solver_stats": {
            sim.name: {"nfev": int(sim.solution.nfev), "njev": int(sim.solution.njev)}
            for sim in [q1, q2, q3, q4]
        },
    }
    (RESULT_DIR / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("[8/8] 生成说明书图表", flush=True)
    create_figures(
        environment,
        radius_history,
        q1,
        q2,
        q3,
        q4,
        appendix4_fixed,
        dry3_s,
        dry4_s,
        safe3_s,
        safe4_s,
        comparison_times_h,
    )
    print("全部数值任务完成。", flush=True)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--no-excel", action="store_true", help="只计算和绘图，不改写附件 3"
    )
    args = parser.parse_args()
    run_all(write_excel=not args.no_excel)


if __name__ == "__main__":
    main()
