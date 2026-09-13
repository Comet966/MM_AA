"""按附件3模板格式写出四个结果工作簿。"""

from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

def reset_sheet(ws) -> None:
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)


def style_header(ws, max_col: int) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1][:max_col]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 13


def write_two_sheet_result(
    path: Path,
    times_s: np.ndarray,
    distances_cm: np.ndarray,
    temperature_c: np.ndarray,
    moisture: np.ndarray,
) -> None:
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = "温度"
        wb.create_sheet("水分浓度")
    mapping = {"温度": temperature_c, "水分浓度": moisture}
    for sheet_name, values in mapping.items():
        ws = wb[sheet_name]
        reset_sheet(ws)
        ws.cell(1, 1, "时间\\到药材中心的距离")
        for col, distance in enumerate(distances_cm, start=2):
            ws.cell(1, col, float(round(distance, 1)))
        for row, time_s in enumerate(times_s, start=2):
            ws.cell(row, 1, int(round(float(time_s))))
            for col, value in enumerate(values[row - 2], start=2):
                cell = ws.cell(row, col, float(round(value, 4)))
                cell.number_format = "0.0000"
        style_header(ws, distances_cm.size + 1)
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_single_sheet_result(
    path: Path,
    times_s: np.ndarray,
    headers: list[float | str],
    values: np.ndarray,
) -> None:
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = "Sheet1"
    ws = wb.active
    reset_sheet(ws)
    ws.cell(1, 1, "时间\\到药材中心的距离")
    for col, header in enumerate(headers, start=2):
        ws.cell(1, col, header)
    for row, time_s in enumerate(times_s, start=2):
        ws.cell(row, 1, int(round(float(time_s))))
        for col, value in enumerate(values[row - 2], start=2):
            if np.isfinite(value):
                cell = ws.cell(row, col, float(round(value, 4)))
                cell.number_format = "0.0000"
    style_header(ws, len(headers) + 1)
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)

