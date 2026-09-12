"""校验支撑材料的结果文件、时间边界与图表完整性。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
RESULT_DIR = ROOT / "result"
FIGURE_DIR = ROOT / "report" / "figures"

EXPECTED_WORKBOOKS = {
    "result1.xlsx": {"sheets": ["温度", "水分浓度"], "rows": 1801, "cols": 22, "first_time": 1, "last_time": 1800},
    "result2.xlsx": {"sheets": ["温度", "水分浓度"], "rows": 10801, "cols": 22, "first_time": 1, "last_time": 10800},
    "result3.xlsx": {"sheets": None, "rows": 3486, "cols": 22, "first_time": 60, "last_time": 209100},
    "result4.xlsx": {"sheets": None, "rows": 3194, "cols": 23, "first_time": 60, "last_time": 191580},
}

EXPECTED_FIGURES = [
    "drying_threshold.png",
    "environment_extension.png",
    "question1_heatmaps.png",
    "question1_profiles.png",
    "question2_heatmaps.png",
    "question2_profiles.png",
    "question3_heatmaps.png",
    "question4_heatmaps.png",
    "radius_and_comparison.png",
    "shrinkage_effect.png",
]


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def validate_workbook(name: str, spec: dict) -> dict:
    path = RESULT_DIR / name
    check(path.is_file(), f"缺少结果文件：{path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if spec["sheets"] is not None:
            check(wb.sheetnames == spec["sheets"], f"{name} 工作表名称或顺序不符：{wb.sheetnames}")
        for ws in wb.worksheets:
            check(ws.max_row == spec["rows"], f"{name}/{ws.title} 行数应为 {spec['rows']}，实际为 {ws.max_row}")
            check(ws.max_column == spec["cols"], f"{name}/{ws.title} 列数应为 {spec['cols']}，实际为 {ws.max_column}")
            check(ws.cell(2, 1).value == spec["first_time"], f"{name}/{ws.title} 首个时间点错误")
            check(ws.cell(ws.max_row, 1).value == spec["last_time"], f"{name}/{ws.title} 末个时间点错误")
        return {"status": "pass", "sheets": wb.sheetnames, "rows": spec["rows"], "cols": spec["cols"]}
    finally:
        wb.close()


def run_validation(require_sensitivity: bool = True) -> dict:
    summary_path = RESULT_DIR / "summary.json"
    check(summary_path.is_file(), "缺少 result/summary.json")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    threshold = summary["threshold_check"]

    check(threshold["question3_safe_60s_output_raw_max"] < 0.15, "问题三报告时点未严格低于 0.15")
    check(threshold["question3_previous_60s_output_raw_max"] >= 0.15, "问题三报告时点不是首个 60 s 严格达标点")
    check(threshold["question4_safe_60s_output_raw_max"] < 0.15, "问题四报告时点未严格低于 0.15")
    check(threshold["question4_previous_60s_output_raw_max"] >= 0.15, "问题四报告时点不是首个 60 s 严格达标点")

    workbook_checks = {name: validate_workbook(name, spec) for name, spec in EXPECTED_WORKBOOKS.items()}
    missing_figures = [name for name in EXPECTED_FIGURES if not (FIGURE_DIR / name).is_file()]
    check(not missing_figures, f"缺少图表：{missing_figures}")

    sensitivity_path = RESULT_DIR / "parameter_sensitivity.json"
    if require_sensitivity:
        check(sensitivity_path.is_file(), "缺少 result/parameter_sensitivity.json")

    report = {
        "status": "pass",
        "workbooks": workbook_checks,
        "strict_threshold": {
            "question3_time_s": summary["drying_time"]["question3_safe_60s_output_s"],
            "question3_raw_max": threshold["question3_safe_60s_output_raw_max"],
            "question4_time_s": summary["drying_time"]["question4_safe_60s_output_s"],
            "question4_raw_max": threshold["question4_safe_60s_output_raw_max"],
        },
        "figures": {"count": len(EXPECTED_FIGURES), "status": "pass"},
        "sensitivity": "pass" if sensitivity_path.is_file() else "skipped",
    }
    (RESULT_DIR / "validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def main(require_sensitivity: bool | None = None) -> None:
    if require_sensitivity is None:
        parser = argparse.ArgumentParser(description=__doc__)
        parser.add_argument("--allow-missing-sensitivity", action="store_true")
        args = parser.parse_args()
        require_sensitivity = not args.allow_missing_sensitivity
    report = run_validation(require_sensitivity=require_sensitivity)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
